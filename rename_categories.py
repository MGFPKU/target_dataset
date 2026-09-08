"""Rename Target_Category values: drop the trailing "目标" / "target" suffix.

Affects Targets_cn.xlsx and Targets_en.xlsx:
  - data sheets (column H values)
  - Overview sheet (category labels + COUNTIFS formula literals)
  - README sheet (column-description rows)
  - changelog entry + "Last update" date

Run from target_dataset/ with plain python.
"""
import datetime
import shutil
from copy import copy

import openpyxl
from openpyxl.cell.cell import MergedCell

CN_OLD_TO_NEW = {
    "碳减排目标": "碳减排",
    "碳强度目标": "碳强度",
    "碳汇目标": "碳汇",
    "节能目标": "节能",
    "能源效率目标": "能源效率",
    "资源效率目标": "资源效率",
    "可再生能源目标": "可再生能源",
    "电气化目标": "电气化",
    "过渡能源目标": "过渡能源",
    "涉煤目标": "涉煤",
    "生产导向目标": "生产导向",
    "工业转型目标": "工业转型",
    "技术创新目标": "技术创新",
    "污染控制目标": "污染控制",
    "环境质量目标": "环境质量",
    "循环利用目标": "循环利用",
    "绿色交通目标": "绿色交通",
    "气候金融目标": "气候金融",
}

EN_OLD_TO_NEW = {
    "Carbon intensity target": "Carbon intensity",
    "Carbon reduction target": "Carbon reduction",
    "Carbon sink target": "Carbon sink",
    "Climate finance target": "Climate finance",
    "Coal-related target": "Coal-related",
    "Electrification target": "Electrification",
    "Energy efficiency target": "Energy efficiency",
    "Energy saving target": "Energy saving",
    "Environmental quality target": "Environmental quality",
    "Green transportation target": "Green transportation",
    "Industrial transformation target": "Industrial transformation",
    "Industrial transformation targets": "Industrial transformation",
    "Pollution control target": "Pollution control",
    "Production-oriented target": "Production-oriented",
    "Production-oriented targets": "Production-oriented",
    "Recycling target": "Recycling",
    "Renewable energy target": "Renewable energy",
    "Resource efficiency target": "Resource efficiency",
    "Technology innovation target": "Technology innovation",
    "Transitional energy target": "Transitional energy",
}

CHANGELOG = {
    "CN": "2026-09-07：简化目标类别名称，删除“目标”后缀（如“可再生能源目标”→“可再生能源”），中英文同步。",
    "EN": '2026-09-07: Simplified target category names by removing the "target" suffix (e.g. "Renewable energy target" → "Renewable energy"), in sync with the Chinese version.',
}


def rename_values(wb: openpyxl.Workbook, old_to_new: dict[str, str]) -> int:
    changed = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c, MergedCell):
                    continue
                v = c.value
                if not isinstance(v, str):
                    continue
                if v.startswith("="):
                    new_v = v
                    for old, new in old_to_new.items():
                        new_v = new_v.replace(f'"{old}"', f'"{new}"')
                    if new_v != v:
                        c.value = new_v
                        changed += 1
                elif v.strip() in old_to_new:
                    c.value = old_to_new[v.strip()]
                    changed += 1
    return changed


def add_changelog(wb: openpyxl.Workbook, sheet_name: str, text: str) -> None:
    ws = wb[sheet_name]
    # changelog sits right below the "Last update" row (row 5); newest first
    for row in ws.iter_rows(min_col=2, max_col=2):
        if row[0].value == text:
            return  # already present (idempotent re-run)
    ws.insert_rows(6)
    src = ws.cell(row=7, column=2)  # previous changelog entry, shifted down
    dst = ws.cell(row=6, column=2)
    dst.value = text
    dst.font = copy(src.font)
    dst.alignment = copy(src.alignment)
    dst.border = copy(src.border)
    dst.fill = copy(src.fill)
    dst.number_format = src.number_format
    ws.cell(row=5, column=3).value = datetime.datetime(2026, 9, 7)


def run(path: str, old_to_new: dict[str, str], lang: str) -> None:
    backup = path.replace(".xlsx", ".pre_category_rename.bak")
    shutil.copy2(path, backup)
    print(f"backup -> {backup}")

    wb = openpyxl.load_workbook(path)
    changed = rename_values(wb, old_to_new)
    add_changelog(wb, "说明" if lang == "CN" else "README", CHANGELOG[lang])
    wb.save(path)
    print(f"{path}: {changed} cells renamed, changelog added")


if __name__ == "__main__":
    run("Targets_cn.xlsx", CN_OLD_TO_NEW, "CN")
    run("Targets_en.xlsx", EN_OLD_TO_NEW, "EN")
