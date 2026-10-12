# -*- coding: utf-8 -*-
"""FYP2614《循环经济发展"十五五"规划》：写入 2 条目标（循环经济表）+ 来源表登记 + 说明/README 日期更新（CN/EN 两份）。"""
import datetime
import shutil

import openpyxl

CN = 'China_Climate_Target_Tracker_cn.xlsx'
EN = 'China_Climate_Target_Tracker_en.xlsx'

SENTENCE_CN = ('到2030年，覆盖生产、消费、回收、利用全链条的循环经济体系更加完善，主要资源产出率比2025年提高16%左右，'
               '大宗固体废弃物年综合利用量达到45亿吨左右，主要再生资源年循环利用量达到5.1亿吨，资源循环利用产业产值达到8万亿元。')
SENTENCE_EN = ('By 2030, the circular economy system covering the whole chain of production, consumption, recycling and '
               'utilisation will be further improved; the output rate of major resources will increase by about 16% compared '
               'with 2025; the annual comprehensive utilisation of bulk solid waste will reach around 4.5 billion tonnes; the '
               'recycling volume of major renewable resources will reach 510 million tonnes; and the output value of the '
               'resource circular utilisation industry will reach 8 trillion yuan.')

SRC_CN = [None, 'FYP2614', 2026, '循环经济发展“十五五”规划', '15th Five-Year Plan for Circular Economy Development',
          '国家发展和改革委员会', '五年规划', '促进减缓', 'Mitigation-enabling', '十五五', '含目标', 'HTM', '中文',
          'https://www.ndrc.gov.cn/xxgk/jd/jd/202607/t20260702_1406253.html']
SRC_EN = [None, 'FYP2614', 2026, '循环经济发展“十五五”规划', '15th Five-Year Plan for Circular Economy Development',
          'National Development and Reform Commission', '五年规划', 'D.2', 'Mitigation-enabling', '15th FYP', 'TAR', 'HTM', 'CN',
          'https://www.ndrc.gov.cn/xxgk/jd/jd/202607/t20260702_1406253.html']

NEW_DATE = datetime.datetime(2026, 10, 12, 0, 0)

# ---------- CN ----------
wb = openpyxl.load_workbook(CN)

ws = wb['循环经济']
ws.append((2026, '资源循环利用产业产值', '达到', '8万亿元', None, '2030', '新目标', '循环利用', '行业',
           SENTENCE_CN, 'FYP2614.htm', '循环经济'))
ws.append((2026, '单位国内生产总值（GDP）用水量', '下降', '10%以上', '2025', '2030', '新目标', '资源效率', '全社会经济',
           '到2030年，单位GDP用水量比2025年下降10%以上。', 'FYP2614.htm', '循环经济'))

ws = wb['来源']
ws.insert_rows(189)
for col, val in enumerate(SRC_CN, start=1):
    ws.cell(row=189, column=col, value=val)
ws['B5'] = '="战略文件总数："&COUNTA(D9:D515)'

ws = wb['说明']
ws['C5'] = NEW_DATE

wb.save(CN)
print('CN done: 循环经济 rows =', wb['循环经济'].max_row)

# ---------- EN ----------
wb = openpyxl.load_workbook(EN)

ws = wb['Circular Economy']
ws.append((2026, 'output value of resource circular utilisation industry', 'reach', '8 trillion yuan', None, '2030',
           'n', 'Recycling', 'sectors', SENTENCE_EN, 'FYP2614.htm', 'Circular Economy'))
ws.append((2026, 'water consumption per unit of GDP', 'reduce by', 'more than 10%', '2025', '2030',
           'n', 'Resource efficiency', 'whole economy',
           'By 2030, water consumption per unit of GDP will decrease by more than 10% compared with 2025.',
           'FYP2614.htm', 'Circular Economy'))

ws = wb['Sources']
ws.insert_rows(189)
for col, val in enumerate(SRC_EN, start=1):
    ws.cell(row=189, column=col, value=val)
ws['B5'] = '="Total number of strategy documents: "&COUNTA(D9:D515)'

ws = wb['README']
ws['C5'] = NEW_DATE

wb.save(EN)
print('EN done: Circular Economy rows =', wb['Circular Economy'].max_row)

# ---------- Data_sources ----------
shutil.copyfile('ndrc_plan.html', 'Data_sources/Filtering Task Website Source CN/FYP2614.htm')
print('Data_sources: FYP2614.htm saved')
