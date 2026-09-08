"""Final CN/EN consistency check: normalized number/unit comparison with
row pairing within (doc, year) groups."""
import re
import openpyxl
from collections import defaultdict

CN, EN = "Targets_cn.xlsx", "Targets_en.xlsx"

CN_COLUMN_MAP = {
    "公布年份": "Announcement_Year", "指标": "Metric", "方向": "Direction",
    "目标值": "Target_Magnitude", "基线": "Baseline",
    "目标年份/时期": "Target_Year_or_Period", "计数": "Count",
    "目标类别": "Target_Category", "责任主体": "Accountability",
    "政策原文": "Sentence", "文件": "Document", "主题标签": "Topic_Label",
    "子主题标签": "Subtopic_Label",
}

CN_DIGIT = {"一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5, "六": 6,
            "七": 7, "八": 8, "九": 9, "〇": 0, "零": 0}
CN_DIGIT10 = dict(CN_DIGIT); CN_DIGIT10["十"] = 10

FYP_MAP = {"十五五": 15, "十四五": 14, "十三五": 13, "十二五": 12, "十一五": 11,
           "十五": 15, "十四": 14, "十三": 13, "十二": 12, "十一": 11,
           "十": 10, "一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6,
           "七": 7, "八": 8, "九": 9}

def parse_cn_int(s):
    total, cur = 0, 0
    for ch in s:
        if ch in "一二两三四五六七八九":
            cur = CN_DIGIT[ch]
        elif ch == "十":
            total += (cur or 1) * 10; cur = 0
        elif ch == "百":
            total += (cur or 1) * 100; cur = 0
        elif ch == "千":
            total += (cur or 1) * 1000; cur = 0
        elif ch == "万":
            total = (total + cur or 1) * 10000; cur = 0
    return total + cur

def cn_numeral(s):
    m = re.match(r"([一两二三四五六七八九十百千]+)(千万|百万|十万|万亿|亿|万|千|百)", s)
    if not m:
        return None
    return parse_cn_int(m.group(1)) * {"百": 1e2, "千": 1e3, "万": 1e4,
        "亿": 1e8, "万亿": 1e12, "十万": 1e5, "百万": 1e6, "千万": 1e7}[m.group(2)]

FYP_CN_RE = re.compile(
    r"(?<![一两二三四五六七八九十之])(?:十五五|十四五|十三五|十二五|十一五|十五|十四|十三|十二|十一|十|[一两二三四五六七八九])五"
    r"(?=(?:规划|期间|时期|末|初|阶段|目标)|$|[\s\"'“”（）()，,。；;])")

def norm_cn_text(s):
    vals = set()
    # digit ranges: "8—10万亿" / "1.2~1.3万"
    for m in re.finditer(r"(?<![A-Za-z0-9.])(\d+(?:\.\d+)?)[—–~-](\d+(?:\.\d+)?)\s*(万亿|亿|万)?", s):
        mult = {"万": 1e4, "亿": 1e8, "万亿": 1e12}.get(m.group(3) or "", 1)
        vals.add(round(float(m.group(1)) * mult, 6))
        vals.add(round(float(m.group(2)) * mult, 6))
    for m in re.finditer(
        r"(?<![A-Za-z0-9.,])(\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?)(?!\.\d)(?![\d—–~-])(?!月)\s*余?\s*"
        r"(万亿|亿|万|百|%|％|Mt|mt|多?家|多?个|多?户|多?套|多?辆|多?次|多?台|多?座)?", s):
        num = float(m.group(1).replace(",", ""))
        mult = {"百": 1e2, "万": 1e4, "亿": 1e8, "万亿": 1e12, "Mt": 1e6,
                "mt": 1e6, "%": 1, "％": 1, "家": 1, "个": 1, "户": 1, "套": 1,
                "辆": 1, "次": 1, "台": 1, "座": 1, "多家": 1, "多个": 1,
                "多户": 1, "多套": 1, "多辆": 1, "多次": 1, "多台": 1,
                "多座": 1}.get(m.group(2) or "", 1)
        vals.add(round(num * mult, 6))
    for m in re.finditer(r"[一两二三四五六七八九十百千]+(?:千万|百万|十万|万亿|亿|万|千|百)", s):
        v = cn_numeral(m.group())
        if v:
            vals.add(round(v, 6))
    # 百分之五十五 -> 55; 百分之三点六 -> 3.6
    for m in re.finditer(r"百分之([一两二三四五六七八九十百千点]+)", s):
        t = m.group(1)
        if "点" in t:
            ip, fp = t.split("点", 1)
            v = float(parse_cn_int(ip)) + float(parse_cn_int(fp)) / (10 ** len(fp))
        else:
            v = float(parse_cn_int(t))
        vals.add(v)
    # counters: 十五种 / 万种 / 百家 / 千家 (not 一次能源, not 第十一个五年,
    # not 万 in "320万户" — a standalone 万/千/百 counts only when no number precedes)
    for m in re.finditer(
        r"(?<![0-9一两二三四五六七八九十])([一两二三四五六七八九十百千万]+)(?:多)?(?:种|个|家|户|项|类|套|辆|台|座|倍|件)(?!五年)", s):
        vals.add(float(parse_cn_int(m.group(1))))
    if "百公里" in s:
        vals.add(100.0)
    # 万元 -> 10,000 (EN "10,000 yuan"); exclude 万亿元 / 1000万元.
    # Only when the text has another number, so bare metric names like
    # "万元工业增加值用水量" don't add a value the EN metric lacks.
    if re.search(r"(?<![\d亿])万元", s) and re.search(r"\d", s):
        vals.add(10000.0)
    for m in FYP_CN_RE.finditer(s):
        vals.add(f"FYP{FYP_MAP[m.group()]}")
    # 第十二个五年 -> FYP12
    for m in re.finditer(r"第([一两二三四五六七八九十]+)个五年", s):
        vals.add(f"FYP{parse_cn_int(m.group(1))}")
    # X千瓦 -> kW plus W- and GW-scale equivalents (mirrors EN kW/MW/GW rules)
    for m in re.finditer(r"(?<![A-Za-z0-9.,])(\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?)(万亿|亿|万)?千瓦(?!时)", s):
        num = float(m.group(1).replace(",", ""))
        mult = {"万": 1e4, "亿": 1e8, "万亿": 1e12}.get(m.group(2) or "", 1)
        v = num * mult
        vals.add(round(v, 6))
        vals.add(round(v * 1e3, 6))
        vals.add(round(v / 1e6, 6))
    # Chinese-numeral power units: 四千万千瓦 / 五十吉瓦
    for m in re.finditer(r"(?<![0-9一两二三四五六七八九十])([一两二三四五六七八九十百千万]+)千瓦(?!时)", s):
        v = float(parse_cn_int(m.group(1)))
        vals.add(round(v, 6))
        vals.add(round(v * 1e3, 6))
        vals.add(round(v / 1e6, 6))
    # X吉瓦 -> gigawatts (bare value plus W- and kW-scale)
    for m in re.finditer(r"(\d+(?:\.\d+)?)吉瓦", s):
        num = float(m.group(1))
        vals.add(round(num, 6))
        vals.add(round(num * 1e9, 6))
        vals.add(round(num * 1e6, 6))
    for m in re.finditer(r"(?<![0-9一两二三四五六七八九十])([一两二三四五六七八九十百千万]+)吉瓦", s):
        num = float(parse_cn_int(m.group(1)))
        vals.add(round(num, 6))
        vals.add(round(num * 1e9, 6))
        vals.add(round(num * 1e6, 6))
    if "千分之一" in s:
        vals.add(0.001)
        vals.add(1.0)
    for m in re.finditer(r"(\d+(?:\.\d+)?)‰", s):
        v = float(m.group(1))
        vals.add(v)
        vals.add(round(v / 1000, 6))
    for y in years_cn(s):
        vals.discard(float(y))
    return vals

WORD_NUM = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6,
            "seven": 7, "eight": 8, "nine": 9, "ten": 10, "eleven": 11,
            "twelve": 12, "thirteen": 13, "fourteen": 14, "fifteen": 15,
            "sixteen": 16, "seventeen": 17, "eighteen": 18, "nineteen": 19,
            "twenty": 20}

def norm_en_text(s):
    vals = set()
    # ranges: "100-150 million tons" (also en-dash / em-dash / tilde)
    for m in re.finditer(
        r"(?<![A-Za-z0-9.])(\d+(?:\.\d+)?)[—–~-](\d+(?:\.\d+)?)\s*"
        r"(trillion|billion|million|thousand)?",
        s, re.IGNORECASE):
        w = (m.group(3) or "").lower()
        mult = {"thousand": 1e3, "million": 1e6, "billion": 1e9,
                "trillion": 1e12}.get(w, 1)
        vals.add(round(float(m.group(1)) * mult, 6))
        vals.add(round(float(m.group(2)) * mult, 6))
    for m in re.finditer(
        r"(?<![A-Za-z0-9.])(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)(?![\d—–~-])(?!st\b|nd\b|rd\b|th\b)"
        r"\s*(trillion|billion|million|thousand|giga|mega|percent|Mt|mt)?",
        s, re.IGNORECASE):
        num = float(m.group(1).replace(",", ""))
        w = (m.group(2) or "").lower()
        mult = {"thousand": 1e3, "million": 1e6, "billion": 1e9,
                "trillion": 1e12, "giga": 1e9, "mega": 1e6,
                "percent": 1, "mt": 1e6}.get(w, 1)
        vals.add(round(num * mult, 6))
    # word numbers: "one million", "two hundred million"
    for m in re.finditer(
        r"\b(one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|"
        r"thirteen|fourteen|fifteen|sixteen|seventeen|eighteen|nineteen|twenty)\s+"
        r"(hundred|thousand|million|billion|trillion)"
        r"(?:\s+(hundred|thousand|million|billion|trillion))?",
        s, re.IGNORECASE):
        mult = {"hundred": 1e2, "thousand": 1e3, "million": 1e6,
                "billion": 1e9, "trillion": 1e12}[m.group(2).lower()]
        v = WORD_NUM[m.group(1).lower()] * mult
        if m.group(3):
            v *= {"hundred": 1e2, "thousand": 1e3, "million": 1e6,
                  "billion": 1e9, "trillion": 1e12}[m.group(3).lower()]
        vals.add(round(v, 6))
    # standalone small number words: "six major categories"
    for m in re.finditer(
        r"\b(six|seven|eight|nine|ten|eleven|twelve|thirteen|fourteen|"
        r"fifteen|sixteen|seventeen|eighteen|nineteen|twenty)\b",
        s, re.IGNORECASE):
        vals.add(float(WORD_NUM[m.group(1).lower()]))
    m = re.search(r"(\d+)(?:st|nd|rd|th) (?:FYP|Five[ -]Year)", s)
    if m:
        vals.add(f"FYP{int(m.group(1))}")
    # power units: add kW/W/GW equivalents (mirrors CN 千瓦/吉瓦 rules).
    # Optional magnitude word covers "50 million kilowatts".
    for m in re.finditer(
        r"(?<![A-Za-z0-9.])(\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?)\s*"
        r"(trillion|billion|million|thousand)?\s*(?:kilowatts?|kW)(?![- ]?h(?:ours?)?\b)",
        s, re.IGNORECASE):
        v = float(m.group(1).replace(",", "")) * {"thousand": 1e3,
            "million": 1e6, "billion": 1e9, "trillion": 1e12}.get(
            (m.group(2) or "").lower(), 1)
        vals.add(round(v, 6))
        vals.add(round(v * 1e3, 6))
        vals.add(round(v / 1e6, 6))
    for m in re.finditer(
        r"(?<![A-Za-z0-9.])(\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?)\s*"
        r"(trillion|billion|million|thousand)?\s*(?:megawatts?|MW)(?![- ]?h(?:ours?)?\b)",
        s, re.IGNORECASE):
        v = float(m.group(1).replace(",", "")) * {"thousand": 1e3,
            "million": 1e6, "billion": 1e9, "trillion": 1e12}.get(
            (m.group(2) or "").lower(), 1)
        vals.add(round(v * 1e3, 6))
        vals.add(round(v * 1e6, 6))
        vals.add(round(v / 1e3, 6))
    for m in re.finditer(
        r"(?<![A-Za-z0-9.])(\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?)\s*"
        r"(trillion|billion|million|thousand)?\s*(?:gigawatts?|GW)(?![- ]?h(?:ours?)?\b)",
        s, re.IGNORECASE):
        v = float(m.group(1).replace(",", "")) * {"thousand": 1e3,
            "million": 1e6, "billion": 1e9, "trillion": 1e12}.get(
            (m.group(2) or "").lower(), 1)
        vals.add(round(v, 6))
        vals.add(round(v * 1e6, 6))
        vals.add(round(v * 1e9, 6))
    for m in re.finditer(r"(\d+(?:\.\d+)?)\s*per mille", s, re.IGNORECASE):
        v = float(m.group(1))
        vals.add(v)
        vals.add(round(v / 1000, 6))
    for y in years_en(s):
        vals.discard(float(y))
    return vals

YEAR_RE = r"(?<!\d)(?:19|20)\d{2}(?!\.\d)(?!\d)"

def years_cn(s):
    out = set()
    for m in re.finditer(YEAR_RE, s):
        after = s[m.end():m.end() + 2]
        if after[:1] in "万亿千百多吨米克元瓦亩辆件台套座个户家次载度" or after.startswith(("公顷",)):
            continue
        out.add(int(m.group()))
    for m in re.finditer(r"[二〇零一两二三四五六七八九]{4}(?=年)", s):
        out.add(int("".join(str(CN_DIGIT[c]) for c in m.group())))
    return out

COUNTER_WORDS = ("units", "enterprises", "companies", "parks", "factories",
                 "institutions", "vehicles", "cars", "boilers", "sets",
                 "stations", "schools", "programs", "products",
                 "cooperatives", "households", "tons", "tonnes", "hectares",
                 "kilometers", "kilometres", "kwh", "m2", "km", "meters",
                 "metres", "farms", "plants", "pieces", "projects", "hospitals")

def years_en(s):
    out = set()
    for m in re.finditer(YEAR_RE, s):
        after = s[m.end():]
        wm = re.match(r"\s*(million|billion|thousand|trillion|giga|mega|mt)\b", after, re.I)
        if wm:
            continue
        wm2 = re.match(r"\s*([a-z]+)\b", after, re.I)
        if wm2 and wm2.group(1).lower() in COUNTER_WORDS:
            continue
        out.add(int(m.group()))
    return out

def sets_consistent(a, b):
    if a == b:
        return True
    rem_a, rem_b = set(a - b), set(b - a)
    if len(rem_a) > 2 or len(rem_b) > 2:
        return False
    for v in list(rem_a):
        if round(v / 1e3, 6) in b or round(v * 1e3, 6) in b:
            rem_a.discard(v)
    for v in list(rem_b):
        if round(v / 1e3, 6) in a or round(v * 1e3, 6) in a:
            rem_b.discard(v)
    return not rem_a and not rem_b

def fyp_end_years(s, cn):
    out = set()
    if cn:
        for m in FYP_CN_RE.finditer(s):
            out.add(1960 + (FYP_MAP[m.group()] - 1) * 5)
        for m in re.finditer(r"第([一两二三四五六七八九十]+)个五年", s):
            out.add(1960 + (parse_cn_int(m.group(1)) - 1) * 5)
    else:
        for m in re.finditer(r"(\d+)(?:st|nd|rd|th)\s+(?:FYP|Five[ -]Year)", s, re.I):
            out.add(1960 + (int(m.group(1)) - 1) * 5)
    return out

def cell_pair_ok(c, e, why):
    if c == e:
        return True
    if c is None or e is None:
        why.append(f"one side empty: CN={c!r} EN={e!r}"); return False
    if isinstance(c, (int, float)) and isinstance(e, (int, float)):
        if abs(float(c) - float(e)) > 1e-9:
            why.append(f"numeric mismatch: CN={c!r} EN={e!r}"); return False
        return True
    if isinstance(c, (int, float)) and isinstance(e, str):
        if (abs(float(c)) not in {abs(x) for x in norm_en_text(e) if not isinstance(x, str)}
                and abs(float(c)) not in years_en(e) | fyp_end_years(e, False)):
            why.append(f"CN numeric {c!r} missing in EN text {e!r}"); return False
        return True
    if isinstance(e, (int, float)) and isinstance(c, str):
        if (abs(float(e)) not in {abs(x) for x in norm_cn_text(c) if not isinstance(x, str)}
                and abs(float(e)) not in years_cn(c) | fyp_end_years(c, True)):
            why.append(f"EN numeric {e!r} missing in CN text {c!r}"); return False
        return True
    if isinstance(c, str) and isinstance(e, str):
        yc, ye = years_cn(c), years_en(e)
        if yc != ye:
            why.append(f"YEARS differ ({sorted(yc)} vs {sorted(ye)}): CN={c!r} EN={e!r}")
            return False
        sc, se = norm_cn_text(c), norm_en_text(e)
        fc = {x for x in sc if isinstance(x, str)}
        fe = {x for x in se if isinstance(x, str)}
        if fc != fe:
            why.append(f"FYP differ: CN={c!r} EN={e!r}")
            return False
        nc = {x for x in sc if not isinstance(x, str)}
        ne = {x for x in se if not isinstance(x, str)}
        if not sets_consistent(nc, ne):
            why.append(f"values differ: CN={c!r} (CN:{sorted(nc)}) "
                       f"EN={e!r} (EN:{sorted(ne)})")
            return False
        return True
    why.append(f"type mismatch: CN={type(c).__name__}:{c!r} EN={type(e).__name__}:{e!r}")
    return False

def row_tokens(row, cn):
    toks = set()
    for v in row.values():
        if isinstance(v, str):
            if cn:
                toks |= {x for x in norm_cn_text(v) if not isinstance(x, str)}
                toks |= years_cn(v)
            else:
                toks |= {x for x in norm_en_text(v) if not isinstance(x, str)}
                toks |= years_en(v)
    return toks

def mag_tokens(row, cn):
    v = row.get("Target_Magnitude")
    if isinstance(v, str):
        if cn:
            return {x for x in norm_cn_text(v) if not isinstance(x, str)} | years_cn(v)
        return {x for x in norm_en_text(v) if not isinstance(x, str)} | years_en(v)
    if isinstance(v, (int, float)):
        return {float(v)}
    return set()

def sim(a, b):
    if not a or not b:
        return 0.0
    return len(a & b) / max(len(a), len(b))

def pair_groups(g1, g2):
    """Greedy pairing of CN rows g1 with EN rows g2, magnitude-weighted."""
    t1 = [(mag_tokens(r, True), row_tokens(r, True), r) for r in g1]
    t2 = [(mag_tokens(r, False), row_tokens(r, False), r) for r in g2]
    pairs = []
    while t1 and t2:
        best = max(((10 * sim(ma, mb) + sim(a, b), i, j)
                    for i, (ma, a, _) in enumerate(t1)
                    for j, (mb, b, _) in enumerate(t2)), default=None)
        if best is None or best[0] < 1.0:
            break
        _, i, j = best
        pairs.append((t1[i][2], t2[j][2]))
        del t1[i], t2[j]
    return pairs, [r for _, _, r in t1], [r for _, _, r in t2]

def load_rows(ws, cn):
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) if h else "" for h in rows[0]]
    cols = {h: i for i, h in enumerate(header)}
    out = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        if cn:
            std = {CN_COLUMN_MAP.get(h, h): row[i] for h, i in cols.items() if h}
        else:
            std = {h: row[i] for h, i in cols.items() if h}
        out.append(std)
    return out

wb_cn = openpyxl.load_workbook(CN, read_only=True, data_only=True)
wb_en = openpyxl.load_workbook(EN, read_only=True, data_only=True)
pairs = list(zip(wb_cn.worksheets, wb_en.worksheets))

total = 0
for ws_c, ws_e in pairs:
    if not ws_c["A1"].value and not ws_e["A1"].value:
        problems = []
        rc = [r for r in ws_c.iter_rows(values_only=True) if any(v is not None for v in r)]
        re_ = [r for r in ws_e.iter_rows(values_only=True) if any(v is not None for v in r)]
        n = max(len(rc), len(re_))
        for i in range(n):
            row_c = rc[i] if i < len(rc) else ()
            row_e = re_[i] if i < len(re_) else ()
            for j in range(max(len(row_c), len(row_e))):
                c = row_c[j] if j < len(row_c) else None
                e = row_e[j] if j < len(row_e) else None
                if c is None and e is None:
                    continue
                why = []
                if not cell_pair_ok(c, e, why):
                    problems.append((i + 1, openpyxl.utils.get_column_letter(j + 1), c, e, why[0]))
        total += len(problems)
        print(f"\n[{ws_c.title} | {ws_e.title}] CN rows={len(rc)} EN rows={len(re_)} "
              f"problems={len(problems)}")
        if len(re_) > len(rc):
            print(f"  EXTRA EN ROW (last): {[v for v in re_[-1]][:14]}")
        for r, col, c, e, why in problems[:6]:
            print(f"  row {r} col {col}: {why}")
        if len(problems) > 6:
            print(f"  ... and {len(problems)-6} more")
        continue

    cn_rows = load_rows(ws_c, cn=True)
    en_rows = load_rows(ws_e, cn=False)
    cn_groups = defaultdict(list)
    for r in cn_rows:
        cn_groups[(str(r.get("Document") or ""), str(r.get("Announcement_Year") or ""))].append(r)
    en_groups = defaultdict(list)
    for r in en_rows:
        en_groups[(str(r.get("Document") or ""), str(r.get("Announcement_Year") or ""))].append(r)

    problems = []
    unmatched = []
    unpaired = []
    for k in sorted(set(cn_groups) | set(en_groups)):
        g1, g2 = cn_groups.get(k, []), en_groups.get(k, [])
        if not g1 or not g2:
            unmatched.append((k, len(g1), len(g2)))
            continue
        paired, left1, left2 = pair_groups(g1, g2)
        for r in left1 + left2:
            unpaired.append((k, r))
        for r1, r2 in paired:
            for col_name in sorted(set(r1) | set(r2)):
                c, e = r1.get(col_name), r2.get(col_name)
                if c is None and e is None:
                    continue
                why = []
                if not cell_pair_ok(c, e, why):
                    problems.append((k, col_name, c, e, why[0]))

    total += len(problems) + len(unmatched) + len(unpaired)
    print(f"\n[{ws_c.title} | {ws_e.title}] CN={len(cn_rows)} EN={len(en_rows)} "
          f"paired groups={len(cn_groups)} problems={len(problems)} "
          f"unmatched_groups={len(unmatched)} unpaired_rows={len(unpaired)}")
    for k, n1, n2 in unmatched[:4]:
        print(f"  GROUP ONLY ONE SIDE: {k} CN={n1} EN={n2}")
    for k, r in unpaired:
        print(f"  UNPAIRED row: {k}: {r.get('Metric')!r} mag={r.get('Target_Magnitude')!r}")
    for k, col, c, e, why in problems:
        print(f"  {k} col={col}: {why}")

print(f"\n=== TOTAL problems: {total} ===")
wb_cn.close(); wb_en.close()
