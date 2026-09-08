# -*- coding: utf-8 -*-
"""Fix 方向/目标值 wording to match 政策原文 exactly in Targets_cn.xlsx (2026-09-08 round 3b).

Round-3b catches rows the round-3 classifier missed (year-digit collisions, Chinese
numerals, non-THRESH verbs) plus value-suffix alignments. All rewrites are verbatim
from the 政策原文 column.

Usage:
  PYTHONIOENCODING=utf-8 python fix_direction3b_cn.py --check   # dry run
  PYTHONIOENCODING=utf-8 python fix_direction3b_cn.py --apply   # backup + fix + save
"""
import sys, shutil
import openpyxl

CN_FILE = 'Targets_cn.xlsx'

# sheet -> [(excel_row, old_dir, new_dir, old_val, new_val)]
# None = column unchanged (still asserted in check())
DIR3B_FIXES = {
    '温室气体排放': [
        (27, '降低', '下降', '3.9%以上', None),
        (46, '降低', '下降', '3.5%', '3.5%以上'),
        (69, '降低', '下降', '20%', None),
        (74, '降低', '下降', '20%', None),
        (110, '达到', '减少', '5亿吨二氧化碳', '约5亿吨二氧化碳'),
    ],
    '能源|电力': [
        (4, '达到', '建成', '20万公里', None),
        (8, '降低', '下降', '比2015年下降5%左右', '5%左右'),
        (14, '降低', '削减', '1亿吨', None),
        (34, '超过', '达到', '95%', '95%以上'),
        (49, '降低', '下降', '较2015年下降10%', '10%'),
        (61, '降低', '下降', '15%', None),
        (64, '降低', '下降', '15%', None),
        (68, '降低', '下降', '15%', None),
        (72, '降低', '下降', '20%', None),
        (73, '降低', '下降', '约20%', None),
        (75, '降低', '下降', '20%', '20%左右'),
        (76, '降低', '下降', '约20%', None),
        (83, '降低', '下降', '15%', None),
        (104, '达到', None, '2亿吨', '2亿吨左右'),
        (125, '占比达到', None, '约95%', '95%左右'),
        (150, '发展', '建成', '1000个示范村', None),
        (198, '实现', '新增', '50亿吨', '50亿吨左右'),
        (250, '淘汰', None, '2000万千瓦', '超过2000万千瓦'),
        (251, '降低', '下降', '2%', None),
        (319, '提高至', '上升至', '27%', None),
        (320, '占比达到', None, '30%', '30%左右'),
        (344, '投产', '投运', '3.55亿千瓦', None),
        (348, '相当于', None, '2100万吨二氧化硫排放', '减少二氧化硫排放2100万吨'),
        (392, '占比达到', '达到', '10%', None),
        (397, '提高至', '达到', '20%左右', None),
        (400, '提高至', '提升至', '15%', None),
        (401, '达到', '提高至', '15%', None),
        (403, '提高至', '提高到', '20%', '20%左右'),
        (412, '达到', None, '20%', '20%左右'),
        (415, '达到', None, '25%', '约25%左右'),
        (418, '提高至', '达到', '10%', None),
        (423, '提高至', '提高到', '20%左右', None),
        (437, '提高', '提高到', '10%', None),
        (467, '超过', '突破', '1亿千瓦', None),
        (543, '保持', '保持在', '90%左右', None),
    ],
    '工业': [
        (51, '降低', '下降', '20%', None),
        (98, '推广', None, '100项', '超过100项'),
        (138, '占比达到', '超过', '20%', None),
        (150, '改造', '超过', '80%以上', '80%'),
    ],
    '建筑': [
        (28, '完成', None, '3.5亿平方米', '超过3.5亿平方米'),
        (29, '提高', '增长', '2亿平方米以上', None),
        (43, '降低', '下降', '20%', None),
        (51, '提高', '增长', '两千万平方米', '2000万平方米以上'),
    ],
    '交通': [
        (18, '达到', '实现', '2000万吨商品燃油', None),
        (30, '达到', '实现', '1500万吨商品燃油', None),
        (39, '占比达到', '达到', '35%', None),
        (40, '占比达到', '达到', '20%', None),
    ],
    '土地利用、土地利用变化和林业': [
        (8, '实现', '增加', '5000万吨二氧化碳', None),
        (113, '提高至', '跃升至', '165亿立方米', None),
        (124, '修复', '改良', '2000万公顷', None),
        (127, '修复', '改良', '1500万公顷', None),
        (138, '种植', '达到', '700万公顷', None),
        (139, '种植', '达到', '350万公顷', None),
        (141, '种植', '完成', '3000万公顷', None),
        (146, '超过', '不少于', '400公里', None),
        (151, '提高至', '提高到', '55%以上', None),
        (160, '达到', '恢复', '5200万公顷', None),
        (161, '提高', '恢复', '5200万公顷', None),
        (178, '实现', None, '50%', '超过50%'),
        (179, '治理', '达', '1000万公顷以上', None),
        (180, '种植', '达到', '260亿株', '260亿株以上'),
        (181, '种植', '达到', '120亿株', '120亿株以上'),
        (184, '提高', '提高到', '40%以上', None),
        (199, '超过', '不低于', '20,000 公顷', '2万公顷'),
    ],
    '循环经济': [
        (6, '完成', None, '60100 吨/日', '6.01万吨/日'),
        (9, '达到', '替代', '13亿吨', None),
        (24, '超过', '突破', '2000万吨', None),
        (63, '占比达到', '达到', '65%', '65%左右'),
        (64, '节约', '实现', '31亿立方米', None),
        (70, '降低至', '降至', '≤3.2立方米/吨', '3.2立方米/吨及以下'),
        (198, '节约', '实现', '10亿立方米', None),
    ],
    '污染': [
        (6, '降低', '下降', '2%', None),
        (24, '降低', '下降', '2%', None),
        (27, '降低', '下降', '20%', None),
        (40, '降低', '减少', '20%', None),
        (46, '保持稳定', '保持在', None, '2005年水平'),
        (62, '降低', '下降', '2%', None),
        (66, '降低', '下降', '15%', '15%以上'),
        (76, '降低', '下降', '30%', None),
        (77, '降低', '下降', '20%', None),
    ],
}


def check():
    wb = openpyxl.load_workbook(CN_FILE, read_only=True)
    n = 0
    errors = []
    for sheet, fixes in DIR3B_FIXES.items():
        ws = wb[sheet]
        for row, old_dir, new_dir, old_val, new_val in fixes:
            cur_dir = ws.cell(row=row, column=3).value
            cur_val = ws.cell(row=row, column=4).value
            if old_dir is not None and cur_dir != old_dir:
                errors.append(f'{sheet} r{row}: dir {cur_dir!r} != expected {old_dir!r}')
            if old_val is not None and cur_val != old_val:
                errors.append(f'{sheet} r{row}: val {cur_val!r} != expected {old_val!r}')
            if old_val is None and cur_val is not None:
                errors.append(f'{sheet} r{row}: val {cur_val!r} != expected None')
            n += 1
    wb.close()
    for e in errors:
        print('MISMATCH:', e)
    if errors:
        sys.exit(1)
    return n


def main():
    apply = '--apply' in sys.argv
    n = check()
    print(f'{n} direction/value fixes verified')
    if not apply:
        print('DRY RUN OK (use --apply to write)')
        return
    shutil.copy(CN_FILE, CN_FILE.replace('.xlsx', '.pre_dirfix3b.bak'))
    wb = openpyxl.load_workbook(CN_FILE)
    for sheet, fixes in DIR3B_FIXES.items():
        ws = wb[sheet]
        for row, old_dir, new_dir, old_val, new_val in fixes:
            if new_dir is not None:
                ws.cell(row=row, column=3).value = new_dir
            if new_val is not None:
                ws.cell(row=row, column=4).value = new_val
    wb.save(CN_FILE)
    wb.close()
    print(f'{CN_FILE}: applied and saved')


if __name__ == '__main__':
    main()
