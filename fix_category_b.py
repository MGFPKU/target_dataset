# -*- coding: utf-8 -*-
"""Apply B-list category fixes (user-judged one by one, 2026-09-08) to CN+EN.

Decisions: B1 建筑r23→节能; B2 交通r62→能源效率; B3 工业r139/140/164→污染控制;
B4 工业r145→节能; B5 交通r74/75→污染控制; B6 LULUCF r12/13→环境质量;
B7 火灾受害率保持碳汇; B8 养殖废弃物保持现状;
B9 循环经济r172/187→能源效率; B10 能源表充电4行→绿色交通;
B11 新能源汽车保持现状; B12 煤矸石保持资源效率;
B13 能源r493/495/521→涉煤; B14 技术普及率保持能源效率。

Usage:
  PYTHONIOENCODING=utf-8 python fix_category_b.py --check   # dry run
  PYTHONIOENCODING=utf-8 python fix_category_b.py --apply   # backup + fix + save
"""
import sys, shutil
import openpyxl

CN_FILE = 'Targets_cn.xlsx'
EN_FILE = 'Targets_en.xlsx'

CN_FIXES = {
    '建筑': [(23, '节能')],
    '交通': [(62, '能源效率'), (74, '污染控制'), (75, '污染控制')],
    '工业': [(139, '污染控制'), (140, '污染控制'), (164, '污染控制'), (145, '节能')],
    '土地利用、土地利用变化和林业': [(12, '环境质量'), (13, '环境质量')],
    '循环经济': [(172, '能源效率'), (187, '能源效率')],
    '能源|电力': [(489, '绿色交通'), (539, '绿色交通'), (549, '绿色交通'), (550, '绿色交通'),
                (493, '涉煤'), (495, '涉煤'), (521, '涉煤')],
}
EN_FIXES = {
    'Buildings': [(23, 'Energy saving')],
    'Transport': [(62, 'Energy efficiency'), (74, 'Pollution control'), (75, 'Pollution control')],
    'Industry': [(139, 'Pollution control'), (140, 'Pollution control'), (164, 'Pollution control'),
                 (145, 'Energy saving')],
    'LULUCF': [(12, 'Environmental quality'), (13, 'Environmental quality')],
    'Circular Economy': [(172, 'Energy efficiency'), (187, 'Energy efficiency')],
    'Energy|Power': [(489, 'Green transportation'), (539, 'Green transportation'),
                     (549, 'Green transportation'), (550, 'Green transportation'),
                     (493, 'Coal-related'), (495, 'Coal-related'), (521, 'Coal-related')],
}

OLD_CN = {
    '建筑': {23: '能源效率'},
    '交通': {62: '节能', 74: '绿色交通', 75: '绿色交通'},
    '工业': {139: '工业转型', 140: '工业转型', 164: '工业转型', 145: '工业转型'},
    '土地利用、土地利用变化和林业': {12: '碳汇', 13: '碳汇'},
    '循环经济': {172: '循环利用', 187: '循环利用'},
    '能源|电力': {489: '电气化', 539: '电气化', 549: '电气化', 550: '电气化',
                493: '能源效率', 495: '能源效率', 521: '能源效率'},
}
OLD_EN = {
    'Buildings': {23: 'Energy efficiency'},
    'Transport': {62: 'Energy saving', 74: 'Green transportation', 75: 'Green transportation'},
    'Industry': {139: 'Industrial transformation', 140: 'Industrial transformation',
                 164: 'Industrial transformation', 145: 'Industrial transformation'},
    'LULUCF': {12: 'Carbon sink', 13: 'Carbon sink'},
    'Circular Economy': {172: 'Recycling', 187: 'Recycling'},
    'Energy|Power': {489: 'Electrification', 539: 'Electrification', 549: 'Electrification',
                     550: 'Electrification', 493: 'Energy efficiency', 495: 'Energy efficiency',
                     521: 'Energy efficiency'},
}

def check(fname, fixes, oldmap):
    wb = openpyxl.load_workbook(fname, read_only=True)
    n = 0
    for sheet, rows in fixes.items():
        ws = wb[sheet]
        for row, cat in rows:
            old = ws.cell(row=row, column=8).value
            exp = oldmap[sheet][row]
            metric = ws.cell(row=row, column=2).value
            if old != exp:
                print(f'ASSERT FAIL {fname} {sheet} r{row} [{metric}]: got {old!r} want {exp!r}')
                wb.close()
                sys.exit(1)
            n += 1
            print(f'OK {fname} {sheet} r{row} [{str(metric)[:30]}] {old} -> {cat}')
    wb.close()
    return n

def main():
    apply = '--apply' in sys.argv
    n = 0
    n += check(CN_FILE, CN_FIXES, OLD_CN)
    n += check(EN_FILE, EN_FIXES, OLD_EN)
    print(f'{n} fixes verified')
    if not apply:
        print('DRY RUN OK (use --apply to write)')
        return
    for fname, fixes in ((CN_FILE, CN_FIXES), (EN_FILE, EN_FIXES)):
        shutil.copy(fname, fname.replace('.xlsx', '.pre_bfix.bak'))
        wb = openpyxl.load_workbook(fname)
        for sheet, rows in fixes.items():
            ws = wb[sheet]
            for row, cat in rows:
                ws.cell(row=row, column=8).value = cat
        wb.save(fname)
        wb.close()
        print(f'{fname}: applied and saved')

if __name__ == '__main__':
    main()
