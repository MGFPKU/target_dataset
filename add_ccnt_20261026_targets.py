# -*- coding: utf-8 -*-
"""2026-10-26 CCNT 更新入库：轻工纺织产业发展“十五五”规划（FYP2615）2 条目标写入 CN/EN 数据集，
来源表登记 FYP2615（含目标）与 FYP2616（无目标），说明/README 日期更新，Data_sources 增补原文。

CN: 工业 +1 行、能源|电力 +1 行、来源表插 FYP2615/FYP2616、说明表最后更新=2026-10-26
EN: Industry +1 行、Energy|Power +1 行、Sources 插 FYP2615/FYP2616、README Last update=2026-10-26
"""
import shutil
import datetime
import openpyxl

S_GF_CN = '专栏1“十五五”时期轻工纺织产业发展主要预期指标：绿色工厂（家）2025年1650、2030年>2000。'
S_BOILER_CN = '推动食品、纺织（含印染）、造纸等行业全面淘汰35蒸吨/小时及以下燃煤锅炉。'

S_GF_EN = ('Column 1: Main expected indicators for the development of the light and textile industries '
           'during the 15th Five-Year Plan period: number of green factories - 1,650 in 2025, >2,000 in 2030.')
S_BOILER_EN = ('Promote the complete phase-out of coal-fired boilers of 35 steam tons per hour and below '
               'in the food, textile (including printing and dyeing) and papermaking industries.')

CN_SRC_FYP2615 = [None, 'FYP2615', 2026, '轻工纺织产业发展“十五五”规划',
                  '15th Five-Year Plan for the Development of the Light and Textile Industries',
                  '工业和信息化部、商务部、市场监管总局', '五年规划', '明确减缓', 'Explicit mitigation',
                  '十五五', '含目标', 'PDF', '中文',
                  'https://www.miit.gov.cn/zwgk/zcwj/wjfb/tz/art/2026/art_9fb9b35709f941d4805d064092032bd0.html']
CN_SRC_FYP2616 = [None, 'FYP2616', 2026, '新型电池产业发展“十五五”规划',
                  '15th Five-Year Plan for the Development of the New-Type Battery Industry',
                  '工业和信息化部等部门', '五年规划', '经济规划（含气候）', 'Economy planning incl. climate',
                  '十五五', '无目标', 'PDF', '中文',
                  'https://www.miit.gov.cn/zwgk/zcwj/wjfb/tz/art/2026/art_255ea93df118460fa105c46c2f546d9a.html']
EN_SRC_FYP2615 = [None, 'FYP2615', 2026, '轻工纺织产业发展“十五五”规划',
                  '15th Five-Year Plan for the Development of the Light and Textile Industries',
                  'Ministry of Industry and Information Technology, Ministry of Commerce and State Administration for Market Regulation',
                  '五年规划', 'B.2', 'Explicit mitigation', '15th FYP', 'TAR', 'PDF', 'CN',
                  'https://www.miit.gov.cn/zwgk/zcwj/wjfb/tz/art/2026/art_9fb9b35709f941d4805d064092032bd0.html']
EN_SRC_FYP2616 = [None, 'FYP2616', 2026, '新型电池产业发展“十五五”规划',
                  '15th Five-Year Plan for the Development of the New-Type Battery Industry',
                  'Ministry of Industry and Information Technology and other departments',
                  '五年规划', 'B.2', 'Economy planning incl. climate', '15th FYP', 'NO-TAR', 'PDF', 'CN',
                  'https://www.miit.gov.cn/zwgk/zcwj/wjfb/tz/art/2026/art_255ea93df118460fa105c46c2f546d9a.html']

NEW_DATE = datetime.datetime(2026, 10, 26, 0, 0)

# ---------- CN ----------
wb = openpyxl.load_workbook('China_Climate_Target_Tracker_cn.xlsx')
shutil.copy('China_Climate_Target_Tracker_cn.xlsx', 'China_Climate_Target_Tracker_cn.pre_ccnt1026.bak.xlsx')

ws = wb['工业']
ws.append((2026, '轻工纺织行业绿色工厂数量', '超过', '2000家', '1650家（2025）', '2030',
           '新目标', '生产导向', '行业', S_GF_CN, 'FYP2615.pdf', '工业'))

ws = wb['能源|电力']
ws.append((2026, '食品、纺织（含印染）、造纸等行业35蒸吨/小时及以下燃煤锅炉', '淘汰', '全面淘汰', None, '十五五',
           '新目标', '化石能源', '行业', S_BOILER_CN, 'FYP2615.pdf', '能源|电力', '煤电'))

ws = wb['来源']
ws.insert_rows(190)
for col, val in enumerate(CN_SRC_FYP2615, start=1):
    ws.cell(row=190, column=col, value=val)
ws.insert_rows(191)
for col, val in enumerate(CN_SRC_FYP2616, start=1):
    ws.cell(row=191, column=col, value=val)
ws['B5'] = '="战略文件总数："&COUNTA(D9:D517)'

wb['说明']['C5'] = NEW_DATE
wb.save('China_Climate_Target_Tracker_cn.xlsx')
print('CN done: 工业 rows =', wb['工业'].max_row, '| 能源|电力 rows =', wb['能源|电力'].max_row)

# ---------- EN ----------
wb = openpyxl.load_workbook('China_Climate_Target_Tracker_en.xlsx')
shutil.copy('China_Climate_Target_Tracker_en.xlsx', 'China_Climate_Target_Tracker_en.pre_ccnt1026.bak.xlsx')

ws = wb['Industry']
ws.append((2026, 'number of green factories in the light and textile industries', 'exceed', '2,000',
           '1,650 (2025)', '2030', 'n', 'Production-oriented', 'sectors', S_GF_EN, 'FYP2615.pdf', 'Industry'))

ws = wb['Energy|Power']
ws.append((2026, 'coal-fired boilers of 35 steam tons per hour and below in the food, textile '
                 '(including printing and dyeing) and papermaking industries', 'phase out', 'completely',
           None, '15th FYP', 'n', 'Fossil energy', 'sectors', S_BOILER_EN, 'FYP2615.pdf',
           'Energy/Power', 'Coal power'))

ws = wb['Sources']
ws.insert_rows(190)
for col, val in enumerate(EN_SRC_FYP2615, start=1):
    ws.cell(row=190, column=col, value=val)
ws.insert_rows(191)
for col, val in enumerate(EN_SRC_FYP2616, start=1):
    ws.cell(row=191, column=col, value=val)
ws['B5'] = '="Total number of strategy documents: "&COUNTA(D9:D517)'

wb['README']['C5'] = NEW_DATE
wb.save('China_Climate_Target_Tracker_en.xlsx')
print('EN done: Industry rows =', wb['Industry'].max_row, '| Energy|Power rows =', wb['Energy|Power'].max_row)

# ---------- Data_sources ----------
shutil.copyfile(r'..\new_policies_raw\2026-10-26\轻工纺织产业发展十五五规划.pdf',
                'Data_sources/Filtering Task PDF Source CN/FYP2615.pdf')
shutil.copyfile(r'..\new_policies_raw\2026-10-26\p2_textile.html',
                'Data_sources/Filtering Task Website Source CN/FYP2615.htm')
print('Data_sources: FYP2615.pdf + FYP2615.htm saved')
