# -*- coding: utf-8 -*-
"""按 compare_output.txt 报告修复 Targets_cn.xlsx / Targets_en.xlsx 的真实数据差异。
用法: python fix_cn_en_data.py [--apply]   (默认 dry-run 只打印)
"""
import re
import sys

import openpyxl

CN_PATH = 'Targets_cn.xlsx'
EN_PATH = 'Targets_en.xlsx'

# 1-based 列号
C = {'metric': 2, 'mag': 4, 'baseline': 5, 'sentence': 10}

# (文件, 表名, 行号, 列, 方式, 旧值, 新值)
# 方式: set=整格覆盖; replace=子串替换(必须命中); regex=正则替换(必须命中)
FIXES = [
    # ---------- GHG Emissions (EN) ----------
    ('en', 'GHG Emissions', 42, 'sentence', 'set', None,
     'From 2014 to 2015, the carbon dioxide emissions per unit of GDP will decrease by more than 4% and 3.5% respectively in two years.'),
    ('en', 'GHG Emissions', 2, 'sentence', 'set', None,
     'The energy-saving and carbon reduction transformation of key areas and industries will result in carbon dioxide emissions reduction of about 130 million tons.'),
    ('en', 'GHG Emissions', 86, 'sentence', 'set', None,
     'From 2024 to 2025, the energy-saving and carbon-reduction transformation of the steel industry will result in energy savings of about 20 million tons of standard coal and reduce carbon dioxide emissions by about 53 million tons.'),
    ('en', 'GHG Emissions', 24, 'sentence', 'set', None,
     'Carbon dioxide emissions per unit of GDP will decline by 17 percent by 2015 compared with 2010.'),
    ('en', 'GHG Emissions', 20, 'sentence', 'replace', 'period (2011-2015), China', 'period, China'),
    ('en', 'GHG Emissions', 30, 'sentence', 'replace', 'CO emission per unit of GDP 2 ', 'CO2 emission per unit of GDP '),
    ('en', 'GHG Emissions', 21, 'sentence', 'replace', 'CO emission', 'CO2 emission'),
    ('en', 'GHG Emissions', 22, 'sentence', 'replace', 'CO emission', 'CO2 emission'),
    ('en', 'GHG Emissions', 32, 'sentence', 'replace', 'CO emission', 'CO2 emission'),
    ('en', 'GHG Emissions', 116, 'mag', 'set', None,
     'effectively control trifluoromethane (HFC-23) emissions'),
    ('en', 'GHG Emissions', 35, 'sentence', 'set', None,
     'binding target of slashing carbon intensity by 18 percent during the 14th Five-Year Plan period (2021-2025).'),
    ('en', 'GHG Emissions', 72, 'sentence', 'replace',
     '9.5 percent from 2020 levels', '9.5 percent by 2030 compared with 2020 levels'),
    ('en', 'GHG Emissions', 11, 'mag', 'set', None, 'before 2030'),
    ('en', 'GHG Emissions', 88, 'sentence', 'regex',
     r'approximately 400%\s*\n\s*10000\s+tons', 'approximately 4 million tons'),

    # ---------- Energy|Power (EN) ----------
    ('en', 'Energy|Power', 16, 'sentence', 'set', None,
     'By 2020, the average coal consumption of existing coal-fired power generation units after transformation will be less than 310 grams per kilowatt-hour, of which the average coal consumption of existing units of 600,000 kilowatts and above (except air-cooled units) after transformation will be less than 300 grams per kilowatt-hour.'),
    ('en', 'Energy|Power', 427, 'mag', 'set', None, '10 percent compared with 2015'),
    ('en', 'Energy|Power', 35, 'mag', 'set', None, 'basically eliminated'),
    ('en', 'Energy|Power', 73, 'mag', 'set', None, 'basically eliminated'),
    ('en', 'Energy|Power', 278, 'metric', 'set', None, 'new coal mines (new construction and expansion)'),
    ('en', 'Energy|Power', 278, 'mag', 'set', None, '450 million tons'),
    ('en', 'Energy|Power', 69, 'sentence', 'replace',
     'In 2015, the electricity saving capacity', 'The electricity saving capacity'),
    ('en', 'Energy|Power', 299, 'sentence', 'replace', ' by 2015.', '.'),
    ('en', 'Energy|Power', 434, 'sentence', 'set', None,
     'By 2015, total national energy consumption and electricity consumption will be controlled at around 4 billion tons of standard coal and 6.15 trillion kilowatt-hours respectively.'),
    ('en', 'Energy|Power', 185, 'sentence', 'set', None,
     'During the 13th Five-Year Plan period, about 40 million kilowatts of conventional hydropower will be newly put into operation across the country, and more than 60 million kilowatts will be started, of which about 5 million kilowatts will be small hydropower.'),
    ('en', 'Energy|Power', 209, 'sentence', 'set', None,
     'By 2020, the newly installed capacity of pumped storage power stations will increase by about 17 million kilowatts, reaching about 40 million kilowatts.'),
    ('en', 'Energy|Power', 326, 'sentence', 'set', None,
     'By 2025, China’s primary crude oil processing capacity will be controlled within 1 billion tons, and about 55% of China’s total oil refining capacity should come from large refineries that have a capacity of at least 10 million tons per year.'),
    ('en', 'Energy|Power', 386, 'sentence', 'replace',
     '10 percent by 2010 compared to 2005, the extraction',
     '10 percent by 2010, and the extraction'),
    ('en', 'Energy|Power', 179, 'sentence', 'replace',
     'reach 350 million kilowatts by 2015.', 'reach 350 million kilowatts.'),
    ('en', 'Energy|Power', 470, 'sentence', 'set', None,
     'In 2024, the national non-hydropower renewable energy electricity consumption responsibility weight will reach about 18 percent, and the scale of renewable energy non-electricity utilization will exceed 60 million tons of standard coal.'),
    ('en', 'Energy|Power', 84, 'sentence', 'replace',
     'per 10,000 yuan of GDP will drop to 2.25 tons of standard coal from',
     'per 10,000 yuan of GDP (at constant 1990 prices) will drop to 2.25 tons of standard coal from'),
    ('en', 'Energy|Power', 430, 'mag', 'set', None, '5 percent compared with 2015'),
    ('en', 'Energy|Power', 430, 'sentence', 'replace',
     ' about 5 percent compared with 2015.', ' about 5 percent.'),
    ('en', 'Energy|Power', 34, 'mag', 'set', None, 'basically eliminated'),
    ('en', 'Energy|Power', 61, 'sentence', 'replace', 'By 2010, ', ''),
    ('en', 'Energy|Power', 146, 'sentence', 'replace', 'By 2010, ', ''),
    ('en', 'Energy|Power', 195, 'sentence', 'replace', 'By 2010, ', ''),
    ('en', 'Energy|Power', 91, 'sentence', 'replace', 'reduce CO emissions.', 'reduce CO2 emissions.'),
    # 能源 CN: HL1701 行 237 政策原文错填为装机句，改为与 19.4% 目标一致
    ('cn', '能源|电力', 237, 'sentence', 'set', None,
     '到2020年，力争水电占总发电量比重达到19.4%左右。'),

    # ---------- Industry (EN) ----------
    ('en', 'Industry', 84, 'sentence', 'set', None,
     'By 2020, the annual production and sales will exceed 2 million, and the cumulative production and sales will exceed 5 million electric vehicles.'),

    # ---------- Buildings (EN) ----------
    ('en', 'Buildings', 15, 'sentence', 'set', None,
     'From 2014 to 2015, the energy consumption per unit building area of public institutions across the country will decrease by 2.2 percent annually, and efforts are made to exceed the target of a 12 percent reduction during the 12th Five-Year Plan period.'),
    ('en', 'Buildings', 75, 'sentence', 'set', None,
     'By 2025, new buildings in cities and towns will fully implement green building standards'),

    # ---------- Transport (EN) ----------
    ('en', 'Transport', 32, 'sentence', 'set', None,
     'By 2015, the maximum speed of battery electric passenger vehicles and plug-in hybrid passenger vehicles shall not be less than 100 km/h, and the comprehensive pure electric driving range shall not be less than 150 km for battery electric vehicles and 50 km for plug-in hybrid vehicles, respectively.'),
    ('en', 'Transport', 33, 'sentence', 'set', None,
     'By 2015, the maximum speed of battery electric passenger vehicles and plug-in hybrid passenger vehicles shall not be less than 100 km/h, and the comprehensive pure electric driving range shall not be less than 150 km for battery electric vehicles and 50 km for plug-in hybrid vehicles, respectively.'),
    ('en', 'Transport', 50, 'mag', 'set', None, 'strive to reach the peak'),
    ('en', 'Transport', 51, 'mag', 'set', None, 'basically eliminated'),
    ('en', 'Transport', 77, 'sentence', 'replace',
     ' annually compared with 2020.', ' annually.'),

    # ---------- LULUCF ----------
    # CN 行 140: 政策原文漏写"00"（mag=2200万公顷，EN 为 22 million）
    ('cn', '土地利用、土地利用变化和林业', 140, 'sentence', 'set', None,
     '到2010年，使2200万公顷沙化土地得到有效治理。'),
    ('en', 'LULUCF', 52, 'sentence', 'set', None,
     'The Chinese government announced that, by 2020 the forest area will be increased by 40 million hectares compared with 2005.'),
    ('en', 'LULUCF', 52, 'baseline', 'set', None, 2005),
    ('cn', '土地利用、土地利用变化和林业', 52, 'baseline', 'set', None, 2005),
    ('en', 'LULUCF', 105, 'sentence', 'replace', 'm3 14 relative', 'm3 relative'),
    ('en', 'LULUCF', 79, 'sentence', 'set', None,
     'The Outline of the 14th Five-Year Plan for National Economic and Social Development of the People’s Republic of China and the Long-Range Objectives Through the Year 2035 clearly sets the quantitative target of increasing the forest coverage rate to 24.1 percent by 2025.'),
    ('en', 'LULUCF', 91, 'sentence', 'set', None,
     'At the 2009 United Nations Climate Change Summit, President Hu Jintao made a solemn promise to the world to “vigorously increase forest carbon sinks and strive to increase forest area by 40 million hectares and forest stock by 1.3 billion cubic meters by 2020 compared with 2005.”'),

    # ---------- Circular Economy ----------
    ('cn', '循环经济', 25, 'metric', 'set', None, '再生废纸年产量'),
    ('cn', '循环经济', 25, 'mag', 'set', None, '8000万吨'),
    ('cn', '循环经济', 32, 'metric', 'set', None, '大宗固体废弃物年度利用量'),
    ('cn', '循环经济', 32, 'mag', 'set', None, '约45亿吨'),
    ('en', 'Circular Economy', 27, 'sentence', 'set', None,
     'By 2015, the annual remanufacturing capacity will reach 800,000 engines, 8 million gearboxes, starters, generators, etc., and 200,000 sets of construction machinery, mining machinery, agricultural machinery, etc., and the annual output value of the remanufacturing industry will reach about 50 billion yuan.'),
    ('en', 'Circular Economy', 28, 'sentence', 'set', None,
     'By 2015, the annual remanufacturing capacity will reach 800,000 engines, 8 million gearboxes, starters, generators, etc., and 200,000 sets of construction machinery, mining machinery, agricultural machinery, etc., and the annual output value of the remanufacturing industry will reach about 50 billion yuan.'),
    ('en', 'Circular Economy', 41, 'sentence', 'set', None,
     'By 2015, the total utilization of major renewable resources will reach 266 million tons, the output value will reach 1.2 trillion yuan, and 18 million people will be employed in the recycling industry.'),
    ('en', 'Circular Economy', 118, 'sentence', 'set', None,
     'By 2020, all county towns and key towns in the country will have sewage treatment capacity, and the sewage treatment rates of cities and county towns will reach about 95 percent and 85 percent respectively.'),
    ('en', 'Circular Economy', 180, 'sentence', 'set', None,
     'By 2025, the water consumption per unit of GDP will decrease by about 16% compared with 2020.'),
    ('en', 'Circular Economy', 50, 'sentence', 'set', None,
     'By 2020, the comprehensive utilization rate of livestock and poultry manure and waste in the country will reach more than 75 percent, and the matching rate of manure and waste treatment facilities and equipment in large-scale farms will reach more than 95 percent.'),

    # ---------- Pollution (EN) ----------
    ('en', 'Pollution', 45, 'mag', 'set', None, 'maintained at the 2005 level'),
    ('en', 'Pollution', 47, 'mag', 'set', None, '2020'),
]


def main():
    apply = '--apply' in sys.argv
    wbs = {
        'cn': openpyxl.load_workbook(CN_PATH),
        'en': openpyxl.load_workbook(EN_PATH),
    }
    errors = []
    n_ok = 0

    def cell(wb_key, sheet, row, col_key):
        return wbs[wb_key][sheet].cell(row=row, column=C[col_key])

    for wb_key, sheet, row, col_key, kind, old, new in FIXES:
        c = cell(wb_key, sheet, row, col_key)
        cur = c.value
        tag = f'{wb_key.upper()} {sheet} r{row} {col_key}'
        if kind == 'set':
            new_val = new
        elif kind == 'replace':
            if not isinstance(cur, str) or old not in cur:
                errors.append(f'{tag}: 未命中子串 {old!r}，当前值 {cur!r}')
                continue
            new_val = cur.replace(old, new)
        elif kind == 'regex':
            if not isinstance(cur, str) or not re.search(old, cur):
                errors.append(f'{tag}: 正则未命中 {old!r}，当前值 {cur!r}')
                continue
            new_val = re.sub(old, new, cur)
        else:
            errors.append(f'{tag}: 未知方式 {kind}')
            continue
        print(f'[{tag}]')
        print(f'  - {cur!r}')
        print(f'  + {new_val!r}')
        if apply:
            c.value = new_val
        n_ok += 1

    # Sources: EN 第 496 行 E 列存在空字符串（CN 为全空），清除以对齐行数
    sc = wbs['en']['Sources'].cell(row=496, column=5)
    print(f'[EN Sources r496 baseline] {sc.value!r} -> None')
    if apply:
        sc.value = None
        n_ok += 1
    else:
        if sc.value != '':
            errors.append(f'EN Sources r496 E 当前值 {sc.value!r}，预期空字符串')

    if errors:
        print('\n== 错误 ==')
        for e in errors:
            print(e)
        sys.exit(1)

    if apply:
        wbs['cn'].save(CN_PATH)
        wbs['en'].save(EN_PATH)
        print(f'\n已应用 {n_ok} 处修改并保存。')
    else:
        print(f'\nDry-run 完成：{n_ok} 处修改待应用（加 --apply 实际写入）。')


if __name__ == '__main__':
    main()
