# -*- coding: utf-8 -*-
"""2026-09-28 CCNT 更新入库：智能网联新能源汽车产业发展“十五五”规划（FYP2613）5 条目标写入 CN/EN 数据集。

CN: 交通 +4 行、温室气体排放 +1 行、来源表插 FYP2613、说明表最后更新=2026-09-28
EN: Transport +4 行、GHG Emissions +1 行、Sources 插 FYP2613、README Last update=2026-09-28
"""
import shutil
import datetime
import openpyxl
from openpyxl.utils import get_column_letter
from copy import copy

S_SALES_CN = '到2030年，我国智能网联新能源汽车全产业链优势进一步巩固，进入世界汽车强国行列，国内市场新能源乘用车、商用车在各自领域新车总销量的占比分别达到70%、40%，具备自动驾驶功能的汽车实现规模应用。'
S_FUEL_CN = '关键技术水平有效提升。乘用车平均燃料消耗量达到每百公里3.3升、纯电动乘用车平均电能消耗量达到每百公里11.5千瓦时左右。'
S_PEAK_CN = '汽车产业国民经济支柱地位进一步凸显，在2030年前实现碳达峰，带动供应链加速绿色低碳转型。'

S_SALES_EN = ("By 2030, China's advantages across the entire industrial chain of intelligent connected "
              "new energy vehicles will be further consolidated, and the country will enter the ranks of "
              "the world's automotive powers. In the domestic market, new energy passenger vehicles and "
              "commercial vehicles will account for 70% and 40%, respectively, of total new vehicle sales "
              "in their respective categories, and vehicles with autonomous driving functions will achieve "
              "large-scale application.")
S_FUEL_EN = ("Key technology levels will be effectively improved. The average fuel consumption of "
             "passenger vehicles will reach 3.3 liters per 100 kilometers, and the average electricity "
             "consumption of pure electric passenger vehicles will reach around 11.5 kilowatt-hours per "
             "100 kilometers.")
S_PEAK_EN = ("The role of the automotive industry as a pillar of the national economy will be further "
             "highlighted, and the industry will achieve carbon peaking before 2030, driving the supply "
             "chain to accelerate its green and low-carbon transformation.")

CN_TRANSPORT_ROWS = [
    [2026, '新能源乘用车在乘用车新车总销量的占比', '达到', '70%', None, 2030, '新目标', '绿色交通', '行业', S_SALES_CN, 'FYP2613.htm', '交通'],
    [2026, '新能源商用车在商用车新车总销量的占比', '达到', '40%', None, 2030, '新目标', '绿色交通', '行业', S_SALES_CN, 'FYP2613.htm', '交通'],
    [2026, '乘用车平均燃料消耗量', '达到', '每百公里3.3升', None, 2030, '新目标', '能源效率', '行业', S_FUEL_CN, 'FYP2613.htm', '交通'],
    [2026, '纯电动乘用车平均电能消耗量', '达到', '每百公里11.5千瓦时左右', None, 2030, '新目标', '能源效率', '行业', S_FUEL_CN, 'FYP2613.htm', '交通'],
]
CN_GHG_ROW = [2026, '汽车产业二氧化碳排放达峰', '实现', '2030年前', None, None, '新目标', '碳减排', '行业', S_PEAK_CN, 'FYP2613.htm', '温室气体排放']

EN_TRANSPORT_ROWS = [
    [2026, 'share of new energy passenger vehicles in total new passenger vehicle sales', 'reach', '70%', None, 2030, 'n', 'Green transportation', 'sectors', S_SALES_EN, 'FYP2613.htm', 'Transport'],
    [2026, 'share of new energy commercial vehicles in total new commercial vehicle sales', 'reach', '40%', None, 2030, 'n', 'Green transportation', 'sectors', S_SALES_EN, 'FYP2613.htm', 'Transport'],
    [2026, 'average fuel consumption of passenger vehicles', 'reach', '3.3 liters per 100 kilometers', None, 2030, 'n', 'Energy efficiency', 'sectors', S_FUEL_EN, 'FYP2613.htm', 'Transport'],
    [2026, 'average electricity consumption of pure electric passenger vehicles', 'reach', 'around 11.5 kilowatt-hours per 100 kilometers', None, 2030, 'n', 'Energy efficiency', 'sectors', S_FUEL_EN, 'FYP2613.htm', 'Transport'],
]
EN_GHG_ROW = [2026, 'carbon dioxide emissions peaking in the automotive industry', 'achieve', 'before 2030', None, None, 'n', 'Carbon reduction', 'sectors', S_PEAK_EN, 'FYP2613.htm', 'GHG Emissions']

CN_SOURCE_ROW = ['FYP2613', 2026, '智能网联新能源汽车产业发展“十五五”规划',
                 'The 15th Five-Year Plan for the Development of the Intelligent Connected New Energy Vehicle Industry',
                 '工业和信息化部等部门', '五年规划', '明确减缓', 'Explicit mitigation', '十五五', '含目标',
                 'PDF', '中文', 'https://www.miit.gov.cn/jgsj/zbys/wjfb/art/2026/art_6136a59a8ff04922b7c73ddbc575434e.html']
EN_SOURCE_ROW = ['FYP2613', 2026, '智能网联新能源汽车产业发展“十五五”规划',
                 'The 15th Five-Year Plan for the Development of the Intelligent Connected New Energy Vehicle Industry',
                 'Ministry of Industry and Information Technology and other departments', '五年规划', 'B.2, B.4',
                 'Explicit mitigation', '15th FYP', 'TAR', 'PDF', 'CN',
                 'https://www.miit.gov.cn/jgsj/zbys/wjfb/art/2026/art_6136a59a8ff04922b7c73ddbc575434e.html']


def copy_row_style(ws, src_row, dst_row, ncols):
    for c in range(1, ncols + 1):
        s = ws.cell(row=src_row, column=c)
        d = ws.cell(row=dst_row, column=c)
        d.font = copy(s.font)
        d.border = copy(s.border)
        d.fill = copy(s.fill)
        d.alignment = copy(s.alignment)
        d.number_format = s.number_format


def append_rows(ws, rows, ncols):
    src = ws.max_row
    start = ws.max_row + 1
    for i, row in enumerate(rows):
        r = start + i
        for j, v in enumerate(row, 1):
            ws.cell(row=r, column=j, value=v)
        copy_row_style(ws, src, r, ncols)


def edit_cn():
    path = 'China_Climate_Target_Tracker_cn.xlsx'
    shutil.copy(path, 'China_Climate_Target_Tracker_cn.pre_ccnt0928.bak.xlsx')
    wb = openpyxl.load_workbook(path)
    append_rows(wb['交通'], CN_TRANSPORT_ROWS, 12)
    append_rows(wb['温室气体排放'], [CN_GHG_ROW], 12)
    ws = wb['来源']
    ws.insert_rows(188)
    for j, v in enumerate(CN_SOURCE_ROW, 2):
        ws.cell(row=188, column=j, value=v)
    copy_row_style(ws, 187, 188, 14)
    ws.cell(row=5, column=2).value = '="战略文件总数："&COUNTA(D9:D514)'
    wb['说明'].cell(row=5, column=3).value = datetime.datetime(2026, 9, 28)
    wb.save(path)
    print('CN saved')


def edit_en():
    path = 'China_Climate_Target_Tracker_en.xlsx'
    shutil.copy(path, 'China_Climate_Target_Tracker_en.pre_ccnt0928.bak.xlsx')
    wb = openpyxl.load_workbook(path)
    append_rows(wb['Transport'], EN_TRANSPORT_ROWS, 12)
    append_rows(wb['GHG Emissions'], [EN_GHG_ROW], 12)
    ws = wb['Sources']
    ws.insert_rows(188)
    for j, v in enumerate(EN_SOURCE_ROW, 2):
        ws.cell(row=188, column=j, value=v)
    copy_row_style(ws, 187, 188, 14)
    ws.cell(row=5, column=2).value = '="Total number of strategy documents: "&COUNTA(D9:D514)'
    wb['README'].cell(row=5, column=3).value = datetime.datetime(2026, 9, 28)
    wb.save(path)
    print('EN saved')


if __name__ == '__main__':
    edit_cn()
    edit_en()
